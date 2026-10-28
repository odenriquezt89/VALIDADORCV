import io
import os
import re
import json
import html
import sqlite3
import smtplib
import bcrypt
import pandas as pd
from datetime import datetime
from typing import Optional
from pypdf import PdfReader
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.application import MIMEApplication
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# Constantes de evaluación
# ---------------------------------------------------------------------------
TARJETA_KEYWORDS = ["tarjeta profesional", "registro medico", "registro médico",
                    "matricula", "matrícula", "rethus", "re-thus"]
LIDERAZGO_KEYWORDS = ["coordinador", "lider", "líder", "gerente", "director", "jefe", "auditor"]
METRICS_REGEX = r'(\d+\s*%|\$\s*\d+|\d+\s*pacientes|\d+\s*casos|\d+\s*usuarios|\d+\s*auditorías|\d+\s*indicadores)'
MESES_REGEX = r'\b(enero|febrero|marzo|abril|mayo|junio|julio|agosto|septiembre|octubre|noviembre|diciembre)\b'
DEFAULT_DB = "data/talent_validator.db"


# ---------------------------------------------------------------------------
# Base de datos
# ---------------------------------------------------------------------------
def get_db_path(secrets_dict: Optional[dict] = None) -> str:
    if os.getenv("DATABASE_PATH"):
        return os.getenv("DATABASE_PATH")
    try:
        if secrets_dict and "database" in secrets_dict:
            return secrets_dict["database"].get("path", DEFAULT_DB)
    except Exception:
        pass
    return DEFAULT_DB


def get_db_connection(db_path: str = DEFAULT_DB) -> sqlite3.Connection:
    db_dir = os.path.dirname(db_path)
    if db_dir:
        os.makedirs(db_dir, exist_ok=True)
    conn = sqlite3.connect(db_path, check_same_thread=False)
    conn.row_factory = sqlite3.Row
    return conn


def init_db(db_path: str = DEFAULT_DB, default_admin_pass: str = "AdminPass2026!"):
    conn = get_db_connection(db_path)
    c = conn.cursor()
    c.execute("""
        CREATE TABLE IF NOT EXISTS usuarios (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT UNIQUE NOT NULL,
            password_hash BLOB NOT NULL,
            rol TEXT NOT NULL,
            activo INTEGER DEFAULT 1,
            plan TEXT DEFAULT 'starter',
            limite_cvs_mes INTEGER DEFAULT 100,
            api_key TEXT UNIQUE,
            fecha_creacion TEXT
        )
    """)
    c.execute("""
        CREATE TABLE IF NOT EXISTS convocatorias (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            codigo TEXT UNIQUE NOT NULL,
            nombre_puesto TEXT NOT NULL,
            meses_experiencia INTEGER DEFAULT 24,
            requiere_tarjeta INTEGER DEFAULT 1,
            skills_json TEXT,
            fecha_creacion TEXT,
            estado TEXT DEFAULT 'ACTIVA'
        )
    """)
    c.execute("""
        CREATE TABLE IF NOT EXISTS candidatos_historico (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            vacante_codigo TEXT NOT NULL,
            nombre TEXT NOT NULL,
            email TEXT,
            telefono TEXT,
            estado_knockout TEXT,
            score_cv REAL,
            score_entrevista REAL,
            score_final REAL,
            detalles_json TEXT,
            seleccionado INTEGER DEFAULT 0,
            fecha_evaluacion TEXT
        )
    """)
    c.execute("CREATE INDEX IF NOT EXISTS idx_cand_vac ON candidatos_historico(vacante_codigo);")
    c.execute("CREATE INDEX IF NOT EXISTS idx_cand_fec ON candidatos_historico(fecha_evaluacion);")

    c.execute("SELECT id FROM usuarios WHERE username = 'admin'")
    if not c.fetchone():
        hashed = bcrypt.hashpw(default_admin_pass.encode('utf-8'), bcrypt.gensalt())
        c.execute("""
            INSERT INTO usuarios (username, password_hash, rol, activo, plan, api_key, fecha_creacion)
            VALUES (?, ?, 'SuperAdmin', 1, 'enterprise', 'key_live_admin_master_2026', ?)
        """, ('admin', hashed, datetime.now().isoformat()))
    conn.commit()
    conn.close()


def save_convocatoria_db(codigo: str, nombre_puesto: str, meses_exp: int, requiere_tarjeta: bool,
                         skills: list, db_path: str = DEFAULT_DB):
    conn = get_db_connection(db_path)
    conn.execute("""
        INSERT INTO convocatorias (codigo, nombre_puesto, meses_experiencia, requiere_tarjeta, skills_json, fecha_creacion, estado)
        VALUES (?, ?, ?, ?, ?, ?, 'ACTIVA')
        ON CONFLICT(codigo) DO UPDATE SET
            nombre_puesto=excluded.nombre_puesto,
            meses_experiencia=excluded.meses_experiencia,
            requiere_tarjeta=excluded.requiere_tarjeta,
            skills_json=excluded.skills_json
    """, (codigo, nombre_puesto, int(meses_exp), 1 if requiere_tarjeta else 0,
          json.dumps(skills, ensure_ascii=False), datetime.now().isoformat()))
    conn.commit()
    conn.close()


def get_convocatorias_db(db_path: str = DEFAULT_DB) -> list:
    conn = get_db_connection(db_path)
    rows = conn.execute("SELECT * FROM convocatorias ORDER BY id DESC").fetchall()
    conn.close()
    return [dict(r) for r in rows]


def save_convocatoria_results(vacante: dict, resumen: list, elegido: Optional[str],
                              correo: Optional[str], db_path: str = DEFAULT_DB) -> int:
    """Guarda (o reemplaza) los resultados de una convocatoria. Evita duplicados."""
    codigo = vacante["codigo_vacante"]
    save_convocatoria_db(codigo, vacante["nombre_puesto"], vacante["meses_experiencia"],
                         vacante["requiere_tarjeta"], vacante["skills_obligatorias"], db_path)
    conn = get_db_connection(db_path)
    conn.execute("DELETE FROM candidatos_historico WHERE vacante_codigo = ?", (codigo,))
    ahora = datetime.now().isoformat()
    for c in resumen:
        es_elegido = c["nombre"] == elegido
        detalles = {
            "m2": c.get("m2", 0.0), "m3": c.get("m3", 0.0), "m4": c.get("m4", 0.0), "m5": c.get("m5", 0.0),
            "fortalezas": c.get("fortalezas", []),
            "debilidades": c.get("debilidades", []),
            "motivo_descarte": c.get("motivo_descarte"),
            "validado_manualmente": c.get("validado_manualmente", False),
            "tarjeta_detectada": c.get("tarjeta_detectada", False),
            "meses_detectados": c.get("meses_detectados", 0),
            "meses_acreditados": c.get("meses_acreditados"),
            "notas_auditor": c.get("notas_auditor", ""),
            "texto_cv": c.get("texto_cv", ""),
            "entrevista_preguntas": c.get("entrevista_detalle", []),
        }
        conn.execute("""
            INSERT INTO candidatos_historico
            (vacante_codigo, nombre, email, estado_knockout, score_cv, score_entrevista, score_final,
             detalles_json, seleccionado, fecha_evaluacion)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (codigo, c["nombre"], correo if es_elegido and correo else None, c["knockout"],
              c["score_cv"], c.get("score_entrevista", 0.0), c.get("score_final", 0.0),
              json.dumps(detalles, ensure_ascii=False), 1 if es_elegido else 0, ahora))
    conn.execute("UPDATE convocatorias SET estado = 'CERRADA' WHERE codigo = ?", (codigo,))
    conn.commit()
    conn.close()
    return len(resumen)


def load_convocatoria(codigo: str, db_path: str = DEFAULT_DB) -> Optional[dict]:
    """Recupera una convocatoria con sus aspirantes y entrevistas."""
    conn = get_db_connection(db_path)
    row = conn.execute("SELECT * FROM convocatorias WHERE codigo = ?", (codigo,)).fetchone()
    if not row:
        conn.close()
        return None
    cand_rows = conn.execute(
        "SELECT * FROM candidatos_historico WHERE vacante_codigo = ? ORDER BY id ASC", (codigo,)
    ).fetchall()
    conn.close()

    vacante = {
        "codigo_vacante": row["codigo"],
        "nombre_puesto": row["nombre_puesto"],
        "meses_experiencia": row["meses_experiencia"],
        "requiere_tarjeta": bool(row["requiere_tarjeta"]),
        "skills_obligatorias": json.loads(row["skills_json"] or "[]"),
    }
    evaluaciones, entrevistas, detalle = [], {}, {}
    for r in cand_rows:
        dt = json.loads(r["detalles_json"]) if r["detalles_json"] else {}
        evaluaciones.append({
            "nombre": r["nombre"],
            "knockout": r["estado_knockout"],
            "score_cv": r["score_cv"] or 0.0,
            "m2": dt.get("m2", 0.0), "m3": dt.get("m3", 0.0),
            "m4": dt.get("m4", 0.0), "m5": dt.get("m5", 0.0),
            "fortalezas": dt.get("fortalezas", []),
            "debilidades": dt.get("debilidades", []),
            "motivo_descarte": dt.get("motivo_descarte"),
            "texto_cv": dt.get("texto_cv", ""),
            "validado_manualmente": dt.get("validado_manualmente", False),
            "tarjeta_detectada": dt.get("tarjeta_detectada", False),
            "meses_detectados": dt.get("meses_detectados", 0),
            "meses_acreditados": dt.get("meses_acreditados"),
            "notas_auditor": dt.get("notas_auditor", ""),
        })
        if r["score_entrevista"]:
            entrevistas[r["nombre"]] = r["score_entrevista"]
        if dt.get("entrevista_preguntas"):
            detalle[r["nombre"]] = dt["entrevista_preguntas"]
    evaluaciones.sort(key=lambda x: (x["knockout"] == "APROBADO", x["score_cv"]), reverse=True)
    return {"vacante": vacante, "evaluaciones": evaluaciones,
            "entrevistas": entrevistas, "entrevistas_detalle": detalle}


# ---------------------------------------------------------------------------
# Extracción y parsing
# ---------------------------------------------------------------------------
def extract_text_from_pdf_bytes(pdf_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(pdf_bytes))
    text = ""
    for page in reader.pages:
        pt = page.extract_text()
        if pt:
            text += pt + "\n"
    return text


def parse_requisicion_text(text: str) -> dict:
    data = {
        "codigo_vacante": "CONV-" + datetime.now().strftime("%Y%m%d-%H%M"),
        "nombre_puesto": "Puesto Requerido",
        "meses_experiencia": 24,
        "requiere_tarjeta": True,
        "skills_obligatorias": []
    }
    for line in [l.strip() for l in text.split("\n") if l.strip()]:
        l_low = line.lower()
        if "cargo:" in l_low or "puesto:" in l_low:
            data["nombre_puesto"] = line.split(":", 1)[-1].strip()
        elif "experiencia:" in l_low:
            nums = re.findall(r'\d+', line)
            if nums:
                val = int(nums[0])
                data["meses_experiencia"] = val * 12 if ("año" in l_low or val < 10) else val
        elif "tarjeta profesional:" in l_low:
            valor = l_low.split(":", 1)[-1]
            negativo = re.search(r'\bno\b', valor) is not None
            afirmativo = re.search(r'\b(si|sí|requiere|obligatori[ao])\b', valor) is not None
            data["requiere_tarjeta"] = afirmativo and not negativo
        elif "habilidades:" in l_low or "competencias:" in l_low or "requisitos:" in l_low:
            raw_s = line.split(":", 1)[-1]
            data["skills_obligatorias"] = [s.strip() for s in raw_s.split(",") if s.strip()]
    return data


def excel_to_requisicion_text(df: pd.DataFrame) -> str:
    """Convierte una hoja Excel en líneas 'Etiqueta: valor' para reutilizar el parser de texto."""
    lines = []
    for _, row in df.iterrows():
        for col, val in row.items():
            if pd.notna(val):
                lines.append(f"{col}: {val}")
    for col in df.columns:
        lines.append(str(col))
    return "\n".join(lines)


def tiene_tarjeta(cv_lower: str) -> bool:
    return any(k in cv_lower for k in TARJETA_KEYWORDS)


def detectar_meses_experiencia(cv_lower: str) -> int:
    matches = re.findall(r'(\d+)\s*(años|año|meses|mes)\s*de\s*experiencia', cv_lower)
    if matches:
        return max(int(n) * (12 if unidad.startswith("añ") else 1) for n, unidad in matches)
    years_found = [int(y) for y in re.findall(r'\b(20[0-2][0-9]|199[0-9])\b', cv_lower)]
    if len(years_found) >= 2:
        return (max(years_found) - min(years_found)) * 12
    return 0


# ---------------------------------------------------------------------------
# Evaluación determinista
# ---------------------------------------------------------------------------
def evaluate_candidate(cv_text: str, vacante: dict) -> dict:
    cv_lower = cv_text.lower()
    passed_ko = True
    motivo_descarte = None

    # Fase 1: Compuerta Knockout (Requisitos Excluyentes)
    if vacante.get("requiere_tarjeta", False) and not tiene_tarjeta(cv_lower):
        passed_ko = False
        motivo_descarte = "Falta mención explícita de Tarjeta Profesional / Matrícula obligatoria."

    exp_months = detectar_meses_experiencia(cv_lower)
    req_months = int(vacante.get("meses_experiencia", 0))
    if passed_ko and exp_months < req_months:
        passed_ko = False
        motivo_descarte = f"Experiencia detectada ({exp_months} meses) menor al umbral exigido ({req_months} meses)."

    if not passed_ko:
        return {
            "knockout": False,
            "motivo_descarte": motivo_descarte,
            "score_cv": 0.0,
            "modulos": {"m2": 0.0, "m3": 0.0, "m4": 0.0, "m5": 0.0},
            "fortalezas": [],
            "debilidades": [motivo_descarte]
        }

    # Fase 2: Ponderación Multicriterio (1 a 100 Puntos)
    # Módulo 2: Fit y Trayectoria (35 pts)
    score_m2 = 35.0 if any(k in cv_lower for k in LIDERAZGO_KEYWORDS) else 25.0

    # Módulo 3: Ajuste Técnico y Hard Skills (30 pts)
    skills = vacante.get("skills_obligatorias", [])
    matched_skills = [s for s in skills if s.lower() in cv_lower]
    ratio_skills = len(matched_skills) / len(skills) if skills else 1.0
    score_m3 = round(ratio_skills * 30.0, 1)

    # Módulo 4: Evidencia de Impacto y Logros (25 pts)
    metrics = re.findall(METRICS_REGEX, cv_lower)
    score_m4 = 25.0 if len(metrics) >= 3 else (18.0 if len(metrics) >= 1 else 10.0)

    # Módulo 5: Calidad Formal y Estructura (10 pts)
    score_m5 = 8.0
    if len(cv_text) > 8500:
        score_m5 -= 3.0
    if re.search(MESES_REGEX, cv_lower):
        score_m5 += 2.0
    score_m5 = min(max(score_m5, 0.0), 10.0)

    total_cv = round(score_m2 + score_m3 + score_m4 + score_m5, 1)

    fortalezas = []
    if score_m3 >= 20:
        fortalezas.append(f"Alineación técnica comprobada ({len(matched_skills)} competencias clave detectadas).")
    if score_m4 >= 18:
        fortalezas.append("Documentación de logros cuantificables y métricas de gestión.")
    if score_m2 == 35.0:
        fortalezas.append("Trayectoria en roles de coordinación, auditoría o liderazgo.")

    debilidades = []
    if score_m3 < 15:
        debilidades.append("Baja presencia de palabras clave técnicas o normativas de la requisición.")
    if score_m4 < 15:
        debilidades.append("Descripción centrada en funciones pasivas sin indicadores de impacto.")

    return {
        "knockout": True,
        "motivo_descarte": None,
        "score_cv": total_cv,
        "modulos": {"m2": score_m2, "m3": score_m3, "m4": score_m4, "m5": score_m5},
        "fortalezas": fortalezas,
        "debilidades": debilidades
    }


def build_candidate_record(nombre: str, cv_text: str, vacante: dict) -> dict:
    """Construye el registro de sesión de un aspirante a partir del texto de su CV."""
    ev = evaluate_candidate(cv_text, vacante)
    cv_lower = cv_text.lower()
    return {
        "nombre": nombre,
        "knockout": "APROBADO" if ev["knockout"] else "DESCARTADO",
        "score_cv": ev["score_cv"],
        "m2": ev["modulos"]["m2"],
        "m3": ev["modulos"]["m3"],
        "m4": ev["modulos"]["m4"],
        "m5": ev["modulos"]["m5"],
        "fortalezas": ev["fortalezas"],
        "debilidades": ev["debilidades"],
        "motivo_descarte": ev["motivo_descarte"],
        "texto_cv": cv_text,
        "validado_manualmente": False,
        "tarjeta_detectada": tiene_tarjeta(cv_lower),
        "meses_detectados": detectar_meses_experiencia(cv_lower),
        "meses_acreditados": None,
        "notas_auditor": "",
    }


def _snippet(text: str, start: int, end: int, before: int = 30, after: int = 45) -> str:
    return text[max(0, start - before):min(len(text), end + after)].replace("\n", " ").strip()


def analyze_cv_for_tables(cv_text: str, vacante: dict, candidato: Optional[dict] = None) -> dict:
    """Organiza la información del CV en 4 tablas temáticas.

    `candidato` puede ser un registro de sesión (claves m2..m5) o un resultado de
    `evaluate_candidate` (clave 'modulos'). Si se omite, se evalúa el texto.
    """
    if candidato is None:
        candidato = evaluate_candidate(cv_text, vacante)
    if "modulos" in candidato:
        modulos = candidato["modulos"]
    else:
        modulos = {k: float(candidato.get(k, 0.0)) for k in ("m2", "m3", "m4", "m5")}

    cv_lower = cv_text.lower()

    # Tabla 1: Datos Generales y Requisitos Clave
    tarjeta_encontrada = "❌ No detectada en el texto"
    for kw in TARJETA_KEYWORDS:
        pos = cv_lower.find(kw)
        if pos != -1:
            tarjeta_encontrada = f"✅ «…{_snippet(cv_text, pos, pos + len(kw), 20, 50)}…»"
            break

    exp_matches = re.findall(r'(\d+\s*(?:años|año|meses|mes)\s*de\s*experiencia)', cv_lower)
    meses = detectar_meses_experiencia(cv_lower)
    exp_texto = (f"{', '.join(exp_matches)} (≈ {meses} meses)" if exp_matches
                 else (f"Inferida por fechas: ≈ {meses} meses" if meses else "Sin mención explícita"))

    years_found = sorted(set(re.findall(r'\b(20[0-2][0-9]|199[0-9])\b', cv_lower)))
    rango_anios = (f"{years_found[0]} – {years_found[-1]}" if len(years_found) >= 2
                   else (years_found[0] if years_found else "No determinados"))

    req_meses = int(vacante.get("meses_experiencia", 0))
    tabla_datos_generales = pd.DataFrame([
        {"Requisito": "Tarjeta Profesional / ReTHUS", "Hallazgo en el documento": tarjeta_encontrada,
         "Exigido": "Obligatoria" if vacante.get("requiere_tarjeta") else "Opcional"},
        {"Requisito": "Experiencia detectada", "Hallazgo en el documento": exp_texto,
         "Exigido": f"Mínimo {req_meses} meses ({'✅ cumple' if meses >= req_meses else '⚠️ no alcanza'})"},
        {"Requisito": "Periodo laboral (años citados)", "Hallazgo en el documento": rango_anios,
         "Exigido": "Cronología consistente"},
        {"Requisito": "Extensión del documento", "Hallazgo en el documento":
            f"{len(cv_text):,} caracteres (~{max(1, len(cv_text) // 1800)} pág.)",
         "Exigido": "Estructura formal sin sobrecarga"},
    ])

    # Tabla 2: Competencias exigidas
    skills_data = []
    for s in vacante.get("skills_obligatorias", []):
        pos = cv_lower.find(s.lower())
        if pos != -1:
            skills_data.append({"Competencia exigida": s, "Estado": "✅ Detectada",
                                "Evidencia": f"«…{_snippet(cv_text, pos, pos + len(s), 25, 45)}…»"})
        else:
            skills_data.append({"Competencia exigida": s, "Estado": "⚠️ No detectada",
                                "Evidencia": "Sin mención explícita en el documento"})
    if not skills_data:
        skills_data = [{"Competencia exigida": "—", "Estado": "N/A",
                        "Evidencia": "La convocatoria no define competencias obligatorias"}]
    tabla_skills = pd.DataFrame(skills_data)

    # Tabla 3: Módulos
    tabla_modulos = pd.DataFrame([
        {"Módulo": "M2 · Trayectoria y liderazgo", "Puntaje": modulos["m2"], "Máximo": 35.0,
         "Criterio": "Cargos de coordinación, dirección, auditoría o jefatura"},
        {"Módulo": "M3 · Hard skills y normativa", "Puntaje": modulos["m3"], "Máximo": 30.0,
         "Criterio": "Proporción de competencias obligatorias presentes"},
        {"Módulo": "M4 · Logros e impacto", "Puntaje": modulos["m4"], "Máximo": 25.0,
         "Criterio": "Indicadores cuantificables, metas, volumen de pacientes"},
        {"Módulo": "M5 · Calidad formal", "Puntaje": modulos["m5"], "Máximo": 10.0,
         "Criterio": "Síntesis, fechas con mes y extensión equilibrada"},
    ])

    # Tabla 4: Logros y métricas
    logros_data = []
    for m in re.finditer(METRICS_REGEX, cv_lower):
        tok = m.group(1)
        tipo = "Porcentaje / eficiencia" if "%" in tok else ("Monetario / presupuesto" if "$" in tok
                                                             else "Volumen asistencial")
        logros_data.append({"Métrica": tok, "Categoría": tipo,
                            "Contexto": f"«…{_snippet(cv_text, *m.span())}…»"})
    if not logros_data:
        logros_data = [{"Métrica": "—", "Categoría": "N/A",
                        "Contexto": "El documento no presenta indicadores numéricos explícitos"}]
    tabla_logros = pd.DataFrame(logros_data)

    return {"datos_generales": tabla_datos_generales, "skills": tabla_skills,
            "modulos": tabla_modulos, "logros": tabla_logros}


def generate_candidate_interview_questions(candidato: dict, vacante: dict) -> list:
    skills = vacante.get("skills_obligatorias", [])
    txt = candidato.get("texto_cv", "").lower()
    puesto = vacante.get("nombre_puesto", "el cargo")

    missing_skills = [s for s in skills if s.lower() not in txt]
    found_skills = [s for s in skills if s.lower() in txt]
    preguntas = []

    # 1. Brecha o dominio técnico
    if missing_skills:
        s = missing_skills[0]
        preguntas.append({
            "tema": f"Brecha técnica: {s}",
            "pregunta": f"En su hoja de vida no se evidencia '{s}'. ¿Cuál es su experiencia práctica y aplicación real de '{s}' en su ejercicio asistencial o administrativo?",
            "foco": f"Validar conocimiento y aplicación real de {s}."
        })
    elif found_skills:
        s = found_skills[0]
        preguntas.append({
            "tema": f"Dominio técnico: {s}",
            "pregunta": f"Usted acredita conocimiento en '{s}'. Describa un caso complejo donde esta competencia fue determinante para la seguridad o la calidad de la atención.",
            "foco": f"Profundizar en la aplicación experta de {s}."
        })
    else:
        preguntas.append({
            "tema": "Marco técnico asistencial",
            "pregunta": f"¿Cómo aplica el marco normativo vigente del sector salud en las funciones diarias del cargo de {puesto}?",
            "foco": "Dominio normativo básico del sector."
        })

    # 2. Liderazgo o trabajo multidisciplinario
    if candidato.get("m2", 0.0) >= 35.0:
        preguntas.append({
            "tema": "Liderazgo de equipos de salud",
            "pregunta": "Usted documenta cargos de coordinación o auditoría. Describa cómo lideró a su equipo durante una contingencia operativa o una auditoría médica rigurosa.",
            "foco": "Liderazgo, manejo de personal y resolución de crisis."
        })
    else:
        preguntas.append({
            "tema": "Trabajo multidisciplinario",
            "pregunta": f"Para el cargo de {puesto}, ¿cómo coordina sus actividades con el equipo multidisciplinario y cómo responde a los requerimientos de supervisión?",
            "foco": "Coordinación, disciplina operativa y trabajo en equipo."
        })

    # 3. Logros
    if candidato.get("m4", 0.0) < 18.0:
        preguntas.append({
            "tema": "Logros cuantificables",
            "pregunta": "Su hoja de vida describe funciones generales. Detalle una meta o mejora cuantificable (tiempos, glosas, oportunidad, adherencia) que haya alcanzado en su última experiencia.",
            "foco": "Comprobar orientación a resultados medibles."
        })
    else:
        preguntas.append({
            "tema": "Indicadores y auditoría",
            "pregunta": "Usted documenta logros cuantificables. ¿Qué indicadores de gestión monitorea con mayor frecuencia y cómo garantiza la confiabilidad de la información?",
            "foco": "Metodología de auditoría, control de gestión e indicadores."
        })

    # 4. Resolución de problemas
    preguntas.append({
        "tema": "Gestión de eventos adversos",
        "pregunta": f"Describa un evento adverso o escenario crítico que haya enfrentado como {puesto} y la secuencia de decisiones que tomó para mitigar el riesgo.",
        "foco": "Criterio clínico/administrativo y toma de decisiones bajo presión."
    })

    # 5. Comunicación y humanización
    preguntas.append({
        "tema": "Comunicación y seguridad del paciente",
        "pregunta": "¿Cómo aborda una discrepancia técnica con un colega o superior, garantizando que el trato al usuario y la seguridad del paciente sigan siendo la prioridad?",
        "foco": "Inteligencia emocional, trato humanizado y resolución asertiva de conflictos."
    })
    return preguntas


# ---------------------------------------------------------------------------
# Reportes
# ---------------------------------------------------------------------------
def _p(text) -> str:
    """Escapa texto libre para Paragraph de ReportLab."""
    return html.escape(str(text if text is not None else ""))


def generate_terna_pdf(vacante: dict, top_3: list, candidato_elegido: str, evaluador_username: str) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    story = []
    styles = getSampleStyleSheet()
    teal, teal_soft, grid = colors.HexColor("#00695C"), colors.HexColor("#E0F2F1"), colors.HexColor("#B2DFDB")

    style_title = ParagraphStyle('DocTitle', parent=styles['Heading1'], fontSize=15, textColor=teal)
    style_h2 = ParagraphStyle('SecH2', parent=styles['Heading2'], fontSize=11, textColor=colors.HexColor("#00897B"),
                              spaceBefore=10, spaceAfter=4)
    style_cell = ParagraphStyle('Cell', parent=styles['Normal'], fontSize=8, leading=10)
    style_bold = ParagraphStyle('CellB', parent=styles['Normal'], fontSize=8, leading=10, fontName='Helvetica-Bold')
    base_style = [('GRID', (0, 0), (-1, -1), 0.5, grid), ('PADDING', (0, 0), (-1, -1), 4),
                  ('VALIGN', (0, 0), (-1, -1), 'TOP')]
    header_style = base_style + [('BACKGROUND', (0, 0), (-1, 0), teal_soft)]

    story.append(Paragraph("INFORME CONSOLIDADO DE SELECCIÓN Y TERNA FINAL", style_title))
    story.append(Paragraph(f"Fecha: {datetime.now().strftime('%d/%m/%Y %H:%M')} | Evaluador: {_p(evaluador_username)}", style_cell))
    story.append(HRFlowable(width="100%", thickness=1.2, color=teal, spaceAfter=8, spaceBefore=4))

    story.append(Paragraph("1. Parámetros de la Convocatoria", style_h2))
    info_vac = [
        [Paragraph("<b>Código:</b>", style_cell), Paragraph(_p(vacante.get("codigo_vacante")), style_cell),
         Paragraph("<b>Cargo:</b>", style_cell), Paragraph(_p(vacante.get("nombre_puesto")), style_cell)],
        [Paragraph("<b>Experiencia:</b>", style_cell), Paragraph(f"{vacante.get('meses_experiencia', 0)} meses", style_cell),
         Paragraph("<b>Tarjeta Prof.:</b>", style_cell), Paragraph("Sí" if vacante.get("requiere_tarjeta") else "No", style_cell)],
        [Paragraph("<b>Competencias:</b>", style_cell),
         Paragraph(_p(", ".join(vacante.get("skills_obligatorias", [])) or "N/A"), style_cell), "", ""]
    ]
    t_vac = Table(info_vac, colWidths=[80, 190, 80, 190])
    t_vac.setStyle(TableStyle(base_style + [('SPAN', (1, 2), (3, 2)),
                                            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#F4FBF9"))]))
    story.append(t_vac)

    story.append(Paragraph("2. Calificación de la Terna Final", style_h2))
    filas_r = [[Paragraph(f"<b>{h}</b>", style_bold) for h in ("#", "Candidato", "CV (60%)", "Entrevista (40%)", "Total")]]
    for idx, c in enumerate(top_3):
        filas_r.append([
            Paragraph(f"#{idx + 1}", style_cell), Paragraph(_p(c["nombre"]), style_bold),
            Paragraph(f"{c['score_cv']}", style_cell), Paragraph(f"{c['score_entrevista']}", style_cell),
            Paragraph(f"<b>{c['score_final']} / 100</b>", style_cell)
        ])
    t_rank = Table(filas_r, colWidths=[35, 225, 90, 95, 95])
    t_rank.setStyle(TableStyle(header_style))
    story.append(t_rank)

    story.append(Paragraph("3. Análisis de Fortalezas y Brechas", style_h2))
    det_filas = [[Paragraph(f"<b>{h}</b>", style_bold) for h in ("Candidato", "Fortalezas", "Brechas")]]
    for c in top_3:
        f_txt = "<br/>".join(f"• {_p(x)}" for x in c["fortalezas"]) or "Sin observaciones."
        d_txt = "<br/>".join(f"• {_p(x)}" for x in c["debilidades"]) or "Sin brechas críticas."
        det_filas.append([Paragraph(_p(c["nombre"]), style_bold), Paragraph(f_txt, style_cell), Paragraph(d_txt, style_cell)])
    t_det = Table(det_filas, colWidths=[130, 205, 205])
    t_det.setStyle(TableStyle(header_style))
    story.append(t_det)

    obs_filas = [[Paragraph(f"<b>{h}</b>", style_bold) for h in ("Candidato", "Pregunta", "Pts", "Observación del evaluador")]]
    for c in top_3:
        for reg in c.get("entrevista_detalle") or []:
            obs_filas.append([Paragraph(_p(c["nombre"]), style_cell),
                              Paragraph(f"P{reg.get('pregunta_num')}: {_p(reg.get('tema'))}", style_cell),
                              Paragraph(str(reg.get("puntaje", 0)), style_cell),
                              Paragraph(_p(reg.get("observacion") or "—"), style_cell)])
    if len(obs_filas) > 1:
        story.append(Paragraph("4. Observaciones de la Entrevista Técnica", style_h2))
        t_obs = Table(obs_filas, colWidths=[100, 150, 30, 260], repeatRows=1)
        t_obs.setStyle(TableStyle(header_style))
        story.append(t_obs)

    story.append(Paragraph(f"{5 if len(obs_filas) > 1 else 4}. Decisión de Selección", style_h2))
    t_dic = Table([[Paragraph(f"<b>Candidato seleccionado:</b> {_p(candidato_elegido)} para el cargo de "
                              f"<b>{_p(vacante.get('nombre_puesto'))}</b>.", style_cell)]], colWidths=[540])
    t_dic.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, -1), teal_soft),
                               ('BOX', (0, 0), (-1, -1), 1, colors.HexColor("#00897B")),
                               ('PADDING', (0, 0), (-1, -1), 6)]))
    story.append(t_dic)

    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()


def generate_excel_historico(df_filtrado: pd.DataFrame) -> bytes:
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df_filtrado.to_excel(writer, index=False, sheet_name='Historico')
        ws = writer.sheets['Historico']
        h_fill = PatternFill(start_color="00897B", end_color="00897B", fill_type="solid")
        h_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        r_font = Font(name="Arial", size=9)
        side = Side(style='thin', color='B2DFDB')
        border = Border(left=side, right=side, top=side, bottom=side)
        for col_num in range(1, len(df_filtrado.columns) + 1):
            c = ws.cell(row=1, column=col_num)
            c.fill, c.font = h_fill, h_font
            c.alignment = Alignment(horizontal="center", vertical="center")
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
            for cell in row:
                cell.font, cell.border = r_font, border
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            ws.column_dimensions[get_column_letter(col[0].column)].width = max(max_len + 4, 12)
    output.seek(0)
    return output.getvalue()


def send_selection_email(smtp_host: str, smtp_port: int, sender_email: str, sender_password: str,
                         recipient_email: str, candidate_name: str, job_title: str, job_code: str, pdf_bytes: bytes) -> dict:
    msg = MIMEMultipart()
    msg['From'] = sender_email
    msg['To'] = recipient_email
    msg['Subject'] = f"Notificación de Selección - {job_code} | {job_title}"
    body = f"""
    <html><body>
      <h3>Notificación Oficial de Selección</h3>
      <p>Estimado/a <b>{html.escape(candidate_name)}</b>,</p>
      <p>Ha sido seleccionado/a en el proceso <b>{html.escape(job_code)}</b> para la posición de
      <b>{html.escape(job_title)}</b>. Se adjunta el acta técnica del proceso.</p>
    </body></html>
    """
    msg.attach(MIMEText(body, 'html'))
    part = MIMEApplication(pdf_bytes, Name=f"Acta_{job_code}.pdf")
    part['Content-Disposition'] = f'attachment; filename="Acta_{job_code}.pdf"'
    msg.attach(part)
    try:
        if smtp_port == 465:
            with smtplib.SMTP_SSL(smtp_host, smtp_port, timeout=20) as s:
                s.login(sender_email, sender_password)
                s.send_message(msg)
        else:
            with smtplib.SMTP(smtp_host, smtp_port, timeout=20) as s:
                s.starttls()
                s.login(sender_email, sender_password)
                s.send_message(msg)
        return {"status": "ok", "message": f"Correo enviado a {recipient_email}."}
    except Exception as e:
        return {"status": "error", "message": str(e)}
